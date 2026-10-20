#!/usr/bin/env python
"""
Buat akun SISWA secara massal dari daftar presensi kelas.

═══════════════════════════════════════════════════════════════════════════
MENGAPA MEMAKAI RPC `teacher_create_student`, BUKAN ADMIN API
═══════════════════════════════════════════════════════════════════════════

  Akun siswa TIDAK dibuat lewat Admin API, meskipun skrip ini bisa saja
  memintanya. Alasannya: kolom `classroom_code` pada `profiles` DIKUNCI
  trigger, sehingga siswa yang dibuat lewat Admin API akan punya
  `classroom_code` kosong — dan ia TIDAK AKAN MUNCUL di Posko Guru.

  Fungsi `teacher_create_student` melakukan semuanya dengan benar dalam satu
  transaksi: membuat pengguna Auth, menulis profil, dan menautkannya ke kelas
  milik guru yang memanggil. Fungsi itu pula yang dipakai aplikasi, jadi
  hasilnya pasti sama dengan pembuatan akun satu per satu lewat Dashboard Guru.

  Konsekuensinya: skrip ini butuh LOGIN GURU (username + kata sandi), bukan
  kunci service_role. Itu justru lebih aman.

═══════════════════════════════════════════════════════════════════════════
CARA MENJALANKAN
═══════════════════════════════════════════════════════════════════════════

  Lihat dulu username yang akan dibuat (TIDAK menyentuh server):

    python scripts/kelola_akun_siswa.py pratinjau "PRESENSI 8A (1).xlsx"

  Buat akunnya (kata sandi guru akan diminta, tidak tampil di layar):

    python scripts/kelola_akun_siswa.py buat "PRESENSI 8A (1).xlsx" buatyas

  Periksa siswa yang sudah terdaftar di kelas guru itu:

    python scripts/kelola_akun_siswa.py periksa buatyas

  Kata sandi siswa bawaan: 123456   (bisa diubah dengan --sandi)
  Kode kelas biasanya terdeteksi otomatis bila guru hanya punya satu kelas;
  bila punya lebih dari satu, sebutkan: --kelas RESQ-XXXX123456

═══════════════════════════════════════════════════════════════════════════
ATURAN YANG DIPAKSA SERVER (dan dipatuhi skrip ini)
═══════════════════════════════════════════════════════════════════════════

  - username  : ^[a-z0-9._-]{3,30}$   (huruf kecil, angka, titik, _ , -)
  - kata sandi: minimal 6 karakter
  - nomor absen: 1–3 digit, UNIK di dalam satu kelas
  - username  : unik di seluruh aplikasi

  Username dibuat dari nama depan siswa. Bila bentrok, ditambah angka
  (mis. `muhammad` lalu `muhammad2`). Daftar lengkapnya bisa dilihat lewat
  perintah `pratinjau`, jadi kesalahan penamaan ketahuan SEBELUM dibuat.
"""
from __future__ import annotations

import getpass
import json
import os
import re
import sys
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

AKAR = Path(__file__).resolve().parent.parent
SANDI_BAWAAN = "123456"


# ── Konfigurasi proyek ────────────────────────────────────────────────────

def _dari_env(nama: str) -> str:
    nilai = os.environ.get(nama, "").strip()
    if nilai:
        return nilai
    env = AKAR / ".env"
    if env.exists():
        for baris in env.read_text(encoding="utf-8", errors="ignore").split("\n"):
            if baris.strip().startswith(nama + "="):
                return baris.split("=", 1)[1].strip().strip('"').strip("'")
    return ""


ALAMAT = _dari_env("VITE_SUPABASE_URL").rstrip("/")
KUNCI_ANON = _dari_env("VITE_SUPABASE_ANON_KEY")


def pastikan_konfigurasi() -> None:
    if not ALAMAT:
        raise SystemExit("  [X] VITE_SUPABASE_URL tidak ditemukan di .env")
    if not KUNCI_ANON:
        raise SystemExit("  [X] VITE_SUPABASE_ANON_KEY tidak ditemukan di .env")


# ── HTTP ──────────────────────────────────────────────────────────────────

def _minta(jalur: str, metode: str = "GET", isi=None, token: str | None = None):
    """Panggil REST/Auth Supabase. Mengembalikan (status, hasil)."""
    data = json.dumps(isi).encode() if isi is not None else None
    req = urllib.request.Request(f"{ALAMAT}{jalur}", method=metode, data=data)
    req.add_header("apikey", KUNCI_ANON)
    req.add_header("Authorization", f"Bearer {token or KUNCI_ANON}")
    req.add_header("Content-Type", "application/json")
    try:
        with urllib.request.urlopen(req, timeout=60) as r:
            teks = r.read().decode("utf-8", "ignore")
            return r.status, (json.loads(teks) if teks.strip() else None)
    except urllib.error.HTTPError as e:
        teks = e.read().decode("utf-8", "ignore")
        try:
            return e.code, json.loads(teks) if teks.strip() else None
        except json.JSONDecodeError:
            return e.code, {"message": teks[:300]}
    except urllib.error.URLError as e:
        raise SystemExit(f"  [X] Tidak dapat menghubungi server: {e.reason}")


def masuk_guru(username: str) -> tuple[str, str]:
    """Login guru. Mengembalikan (access_token, user_id)."""
    email = f"{username.lower()}@resqbox.local"
    print(f"  Masuk sebagai guru `{username}` ...")
    sandi = os.environ.get("RESQ_GURU_SANDI") or getpass.getpass(
        "  Kata sandi guru (tidak terlihat): "
    )
    status, hasil = _minta(
        "/auth/v1/token?grant_type=password",
        "POST",
        {"email": email, "password": sandi},
    )
    if status != 200 or not isinstance(hasil, dict) or "access_token" not in hasil:
        pesan = hasil.get("error_description") or hasil.get("msg") or hasil.get("message") if isinstance(hasil, dict) else hasil
        raise SystemExit(
            f"  [X] Gagal masuk (HTTP {status}): {pesan}\n"
            "      Periksa username dan kata sandi guru. Akun guru dibuat lewat:\n"
            "      python scripts/kelola_akun_guru.py buat <username> \"<Nama>\" \"<Sekolah>\""
        )
    return hasil["access_token"], hasil["user"]["id"]


def profil_saya(token: str) -> dict:
    """Profil pemanggil, termasuk penanda admin."""
    status, hasil = _minta("/rest/v1/rpc/get_my_profile", "POST", {}, token=token)
    if status != 200 or not isinstance(hasil, dict):
        return {}
    return hasil


def kelas_milik_guru(token: str, username_guru: str) -> list[dict]:
    q = urllib.parse.urlencode(
        {"select": "code,name", "teacher_username": f"eq.{username_guru.lower()}"}
    )
    status, hasil = _minta(f"/rest/v1/classrooms?{q}", token=token)
    if status != 200:
        raise SystemExit(f"  [X] Gagal membaca kelas (HTTP {status}): {hasil}")
    return hasil or []


def panggil_buat_siswa(token: str, kode: str, nama: str, absen: str,
                       username: str, sandi: str):
    return _minta(
        "/rest/v1/rpc/teacher_create_student",
        "POST",
        {
            "p_classroom_code": kode,
            "p_name": nama,
            "p_absent_number": absen,
            "p_username": username,
            "p_password": sandi,
        },
        token=token,
    )


# ── Membaca presensi ──────────────────────────────────────────────────────

def baca_presensi(berkas: Path) -> list[dict]:
    """Baca kolom NO dan NAMA SISWA dari berkas presensi (.xlsx)."""
    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError:
        raise SystemExit(
            "  [X] openpyxl belum terpasang.\n"
            "      Pasang dengan: pip install openpyxl"
        )

    if not berkas.exists():
        raise SystemExit(f"  [X] Berkas tidak ditemukan: {berkas}")

    wb = load_workbook(berkas, data_only=True)
    ws = wb.worksheets[0]
    siswa: list[dict] = []
    kolom_no = None
    kolom_nama = None

    for baris in ws.iter_rows(values_only=True):
        sel = ["" if v is None else str(v).strip() for v in baris]
        naik = [s.upper() for s in sel]

        # Cari baris judul kolom
        if "NAMA SISWA" in naik:
            kolom_nama = naik.index("NAMA SISWA")
            kolom_no = naik.index("NO") if "NO" in naik else (kolom_nama - 1)
            continue

        if kolom_nama is None:
            continue

        no = sel[kolom_no] if kolom_no is not None and kolom_no < len(sel) else ""
        nama = sel[kolom_nama] if kolom_nama < len(sel) else ""
        if not nama or not re.fullmatch(r"\d{1,3}", no):
            continue
        nama = re.sub(r"\s+", " ", nama).strip()
        if not re.search(r"[A-Za-z]", nama):
            continue
        siswa.append({"absen": no, "nama": nama})

    if not siswa:
        raise SystemExit(
            "  [X] Tidak ada baris siswa yang terbaca.\n"
            "      Pastikan ada kolom berjudul NO dan NAMA SISWA."
        )
    return siswa


# ── Username ──────────────────────────────────────────────────────────────

def _bersihkan(teks: str) -> str:
    teks = unicodedata.normalize("NFKD", teks)
    teks = "".join(c for c in teks if not unicodedata.combining(c))
    teks = teks.lower().replace("'", "").replace("’", "")
    return re.sub(r"[^a-z0-9]", "", teks)


def bikin_username(nama: str, terpakai: set[str]) -> str:
    """
    Username dari nama depan. Bila bentrok, tambah angka.

    Harus memenuhi ^[a-z0-9._-]{3,30}$ yang dipaksa server.
    """
    bagian = [_bersihkan(b) for b in nama.split()]
    bagian = [b for b in bagian if b]

    calon = bagian[0] if bagian else ""
    if len(calon) < 3:
        calon = "".join(bagian)
    if len(calon) < 3:
        calon = "siswa"
    calon = calon[:28]

    akhir = calon
    n = 2
    while akhir in terpakai:
        akhir = f"{calon}{n}"
        n += 1
    terpakai.add(akhir)
    return akhir


def siapkan_data(siswa: list[dict], sandi: str) -> list[dict]:
    """Lengkapi tiap siswa dengan username & kata sandi."""
    terpakai: set[str] = set()
    hasil = []
    for s in siswa:
        u = bikin_username(s["nama"], terpakai)
        hasil.append({**s, "username": u, "sandi": sandi})
    return hasil


# ── Perintah ──────────────────────────────────────────────────────────────

def perintah_pratinjau(berkas: Path, sandi: str) -> None:
    data = siapkan_data(baca_presensi(berkas), sandi)
    print(f"\n  {len(data)} siswa akan dibuat dengan data berikut:")
    print(f"  {'NO':>3}  {'ABSEN':>5}  {'USERNAME':<20} {'KATA SANDI':<12} NAMA")
    print(f"  {'-'*3}  {'-'*5}  {'-'*20} {'-'*12} {'-'*34}")
    for i, d in enumerate(data, 1):
        print(f"  {i:>3}  {d['absen']:>5}  {d['username']:<20} {d['sandi']:<12} {d['nama']}")
    print("\n  Ini PRATINJAU — belum ada yang dibuat di server.")
    print(f"  Untuk membuatnya: python scripts/kelola_akun_siswa.py buat \"{berkas.name}\" <username_guru>\n")


def perintah_buat(berkas: Path, username_guru: str, kode_kelas: str | None,
                  sandi: str) -> None:
    pastikan_konfigurasi()
    data = siapkan_data(baca_presensi(berkas), sandi)

    print(f"\n  {len(data)} siswa akan dibuatkan akun.")
    token, _ = masuk_guru(username_guru)

    # Tentukan kelas.
    #
    # Pemilik kelas biasanya harus guru yang memanggil. Namun server juga
    # mengizinkan ADMIN menambah murid ke kelas milik guru lain (lihat
    # teacher_create_student: pemeriksaan kepemilikan dilewati bila
    # v_me.is_admin). Tanpa cabang ini, admin yang menyebut --kelas milik
    # orang lain akan ditolak skrip padahal server mengizinkannya.
    kelas = kelas_milik_guru(token, username_guru)
    saya = profil_saya(token)
    admin = bool(saya.get("is_admin")) or str(saya.get("role") or "") == "admin"

    if kode_kelas:
        kode = kode_kelas.upper()
        if not any(k["code"] == kode for k in kelas):
            if admin:
                print(f"  [i] Kelas {kode} bukan milik `{username_guru}`, tetapi akun ini ADMIN")
                print("      sehingga diizinkan server. Kelas akan divalidasi saat pembuatan.")
            else:
                daftar = ", ".join(k["code"] for k in kelas) or "(tidak punya kelas)"
                raise SystemExit(
                    f"  [X] Kelas {kode} bukan milik `{username_guru}`.\n"
                    f"      Kelas miliknya: {daftar}\n"
                    "      Yang dapat menambah ke kelas guru lain hanyalah akun ADMIN."
                )
    elif len(kelas) == 0 and admin:
        raise SystemExit(
            "  [X] Akun ini admin dan tidak punya kelas sendiri.\n"
            "      Sebutkan kelas tujuannya: --kelas RESQ-8A"
        )
    elif len(kelas) == 1:
        kode = kelas[0]["code"]
        print(f"  Kelas terdeteksi: {kelas[0]['name']} ({kode})")
    elif not kelas:
        raise SystemExit(
            "  [X] Guru ini belum punya kelas.\n"
            "      Buat kelas dulu di aplikasi: Posko Guru -> BUAT KELAS BARU."
        )
    else:
        daftar = ", ".join(f"{k['code']} ({k['name']})" for k in kelas)
        raise SystemExit(
            f"  [X] Guru punya lebih dari satu kelas: {daftar}\n"
            "      Sebutkan salah satu dengan --kelas <KODE>"
        )

    print(f"\n  Membuat akun di kelas {kode} ...\n")
    berhasil, sudah_ada, gagal = [], [], []

    for d in data:
        status, hasil = panggil_buat_siswa(
            token, kode, d["nama"], d["absen"], d["username"], d["sandi"]
        )
        pesan = ""
        if isinstance(hasil, dict):
            pesan = str(hasil.get("message") or hasil.get("hint") or "")

        if status in (200, 201):
            berhasil.append(d)
            print(f"  [v] {d['username']:<20} {d['nama']}")
        elif "sudah ada" in pesan.lower() or "sudah terdaftar" in pesan.lower():
            sudah_ada.append({**d, "pesan": pesan})
            print(f"  [=] {d['username']:<20} {d['nama']}  (sudah ada, dilewati)")
        else:
            gagal.append({**d, "pesan": pesan or f"HTTP {status}"})
            print(f"  [X] {d['username']:<20} {d['nama']}  -> {pesan or status}")

    # Simpan daftar akun untuk dibagikan ke siswa
    csv = berkas.with_name(berkas.stem + "-akun.csv")
    with csv.open("w", encoding="utf-8-sig", newline="") as f:
        f.write("NO,ABSEN,NAMA,USERNAME,KATA SANDI,STATUS\n")
        for i, d in enumerate(data, 1):
            if d in berhasil:
                st = "dibuat"
            elif any(x["username"] == d["username"] for x in sudah_ada):
                st = "sudah ada"
            else:
                st = "GAGAL"
            # Nama diapit dan tanda kutip di dalamnya digandakan, supaya nama
            # berkoma tidak menggeser kolom.
            nama_aman = '"' + d['nama'].replace('"', '""') + '"'
            f.write(
                f"{i},{d['absen']},{nama_aman},{d['username']},{d['sandi']},{st}\n"
            )

    print(f"\n  {'='*66}")
    print(f"  Berhasil dibuat : {len(berhasil)}")
    print(f"  Sudah ada       : {len(sudah_ada)}")
    print(f"  Gagal           : {len(gagal)}")
    print(f"  {'='*66}")
    print(f"  Daftar akun disimpan: {csv}")

    if gagal:
        print("\n  Akun yang GAGAL — periksa pesannya:")
        for g in gagal:
            print(f"    - {g['nama']} ({g['username']}): {g['pesan']}")
    if berhasil or sudah_ada:
        print(
            "\n  Sampaikan ke siswa: username + kata sandi."
            f"\n  Kata sandi bawaan: {sandi} — mintalah siswa menggantinya lewat"
            "\n  Profil -> GANTI PASSWORD AKUN SISWA setelah masuk pertama kali."
        )


def perintah_periksa(username_guru: str) -> None:
    pastikan_konfigurasi()
    token, _ = masuk_guru(username_guru)
    kelas = kelas_milik_guru(token, username_guru)
    if not kelas:
        raise SystemExit("  [X] Guru ini belum punya kelas.")

    for k in kelas:
        print(f"\n  Kelas {k['name']} ({k['code']})")
        q = urllib.parse.urlencode(
            {
                "select": "full_name,absent_number,username",
                "classroom_code": f"eq.{k['code']}",
                "order": "absent_number.asc",
            }
        )
        status, hasil = _minta(f"/rest/v1/profiles?{q}", token=token)
        if status != 200:
            print(f"    [X] Gagal membaca profil (HTTP {status})")
            continue
        baris = hasil or []
        if not baris:
            print("    (belum ada siswa)")
            continue
        print(f"    {'ABSEN':>5}  {'USERNAME':<20} NAMA")
        for s in baris:
            print(
                f"    {str(s.get('absent_number') or '-'):>5}  "
                f"{str(s.get('username') or '-'):<20} {s.get('full_name') or '-'}"
            )


# ── Titik masuk ───────────────────────────────────────────────────────────

def main() -> None:
    arg = sys.argv[1:]
    if not arg or arg[0] in ("-h", "--help", "help"):
        print(__doc__)
        return

    perintah = arg[0]
    sisa = arg[1:]

    # Pisahkan opsi (--nama nilai) dari argumen posisi, tanpa mengandalkan
    # sisa.index() yang akan salah bila sebuah nilai muncul dua kali.
    opsi: dict[str, str] = {}
    posisi: list[str] = []
    i = 0
    while i < len(sisa):
        a = sisa[i]
        if a.startswith("--"):
            nilai = sisa[i + 1] if i + 1 < len(sisa) else ""
            opsi[a[2:]] = nilai
            i += 2
        else:
            posisi.append(a)
            i += 1

    sandi = opsi.get("sandi") or SANDI_BAWAAN
    kelas = opsi.get("kelas")

    if perintah == "pratinjau":
        if not posisi:
            raise SystemExit("  [X] Sebutkan berkas presensi (.xlsx)")
        perintah_pratinjau(Path(posisi[0]), sandi)

    elif perintah == "buat":
        if len(posisi) < 2:
            raise SystemExit(
                "  [X] Pemakaian: buat <berkas.xlsx> <username_guru> [--kelas KODE] [--sandi 123456]"
            )
        perintah_buat(Path(posisi[0]), posisi[1], kelas, sandi)

    elif perintah == "periksa":
        perintah_periksa(posisi[0] if posisi else "")

    else:
        raise SystemExit(f"  [X] Perintah tidak dikenal: {perintah}\n      Coba: pratinjau | buat | periksa")


if __name__ == "__main__":
    main()
